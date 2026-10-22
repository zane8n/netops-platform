"""Correlate an operator's own Excel workbooks with Netrune's inventory.

An operator keeps a sheet of switches; Netrune's scans are fresher. This reads
the workbook, finds each sheet's header row, works out which column holds which
field, matches every data row to exactly one switch, and writes Netrune's values
into the cells that are empty or out of date -- leaving every other cell, sheet,
style, and formula as the operator built it.

Nothing here touches a switch or Netrune's database. The CLI layer
(`excel_cli.py`) owns the preview/backup/apply/verify/audit lifecycle; this
module is the pure planning and workbook mechanics so it can be tested without
an App.
"""

from __future__ import annotations

import copy
import datetime as dt
import ipaddress
import re
import time
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Set, Tuple

from netops_platform.identity import _host_label, normalize_mac
from netops_platform.inventory_ingest import _meaningful_model


FIELDS = ("hostname", "ip", "manufacturer", "model", "serial", "mac", "platform", "status", "last_seen")
IDENTITY_FIELDS = ("serial", "mac", "ip", "hostname")
# Status and Last Seen are opt-in (--fields): an asset sheet's "Status" is
# usually lifecycle ("In stock", "Decommissioned"), not reachability.
DEFAULT_FIELDS = ("hostname", "ip", "manufacturer", "model", "serial", "mac", "platform")
FIELD_LABELS = {
    "hostname": "Hostname", "ip": "IP Address", "manufacturer": "Manufacturer", "model": "Model", "serial": "Serial Number",
    "mac": "Base MAC", "platform": "Platform", "status": "Status", "last_seen": "Last Seen",
}

# Normalized header text (lowercase letters and digits only) to field.
HEADER_SYNONYMS: Dict[str, Set[str]] = {
    "hostname": {"hostname", "host", "name", "devicename", "switchname", "switch", "device", "sysname",
                 "systemname", "nodename", "node", "switchhostname", "devicehostname"},
    "ip": {"ip", "ipaddress", "ipaddr", "mgmtip", "managementip", "mgmtipaddress", "managementipaddress",
           "ipv4", "ipv4address", "switchip", "deviceip", "mgmtaddress", "managementaddress"},
    "mac": {"mac", "macaddress", "basemac", "basemacaddress", "basemacaddr", "baseethernetmac",
            "baseethernetmacaddress", "systemmac", "chassismac", "switchmac"},
    "serial": {"serial", "serialnumber", "serialno", "serialnum", "sn", "sno", "systemserial",
               "systemserialnumber", "chassisserial", "motherboardserial", "motherboardserialnumber"},
    "model": {"model", "modelnumber", "modelno", "devicemodel", "switchmodel", "hardwaremodel",
              "hardware", "pid", "productmodel"},
    # A Vendor/Make column wants the maker (Cisco), not netrune's driver label
    # (Cisco IOS); Platform/OS keeps the label.
    "manufacturer": {"manufacturer", "vendor", "make", "brand", "oem", "mfr", "mfg", "devicevendor"},
    "platform": {"platform", "os", "devicetype", "operatingsystem"},
    "status": {"status", "state", "reachability", "onlinestatus", "netrunestatus"},
    "last_seen": {"lastseen", "lastpolled", "lastcontact", "lastupdated", "lastchecked"},
    # Free text, never a device value: only --remarks / --free-ips write it.
    "remarks": {"remarks", "remark", "comments", "comment", "notes", "note"},
}
REMARKS = "remarks"
REMARKS_HEADER = "Remarks"
REMARK_LABELS = {
    "hostname": "name", "ip": "IP", "manufacturer": "vendor", "model": "model", "serial": "serial",
    "mac": "MAC", "platform": "platform", "status": "status", "last_seen": "last seen",
}
# Remarks this module writes about an address or a switch's state. They are
# replaced on every run, so "Free to use" never outlives the address being
# free; anything else in the cell is the operator's and is kept.
OWN_REMARK = re.compile(
    r"^(?:free to use(?: \((?:was .+ until \d{4}-\d{2}-\d{2}|old entry)\))?"
    r"|in use \((?:answers ping|answers on port \d+|switch .+|AP .+|device .+)\)"
    r"|used by .+ until \d{4}-\d{2}-\d{2}|no reply, was .+ until \d{4}-\d{2}-\d{2}"
    r"|allocated, offline|offline|last seen \d{4}-\d{2}-\d{2})$",
    re.IGNORECASE,
)
FREE_OLD_ENTRY = "Free to use (old entry)"
ALLOCATED_OFFLINE = "Allocated, offline"
# An operator remark like these means the address is spoken for even when
# nothing answers on it.
RESERVED_WORDS = re.compile(
    r"reserv|assign|allocat|in use|used by|dhcp|gateway|\bgw\b|hsrp|vrrp|\bvip\b|do not|don'?t|static|spare for",
    re.IGNORECASE,
)

# Parts of a workbook openpyxl does not carry through a load/save round trip.
FEATURES_LOST_ON_SAVE = {
    "xl/charts/": "charts",
    "xl/drawings/": "drawings, shapes, or images",
    "xl/media/": "embedded images",
    "xl/pivotTables/": "pivot tables",
    "xl/slicers/": "slicers",
    "xl/timelines/": "timelines",
    "xl/metadata.xml": "dynamic-array (spill) formulas",
    "xl/externalLinks/": "external workbook links",
}

PLACEHOLDERS = {"", "n/a", "na", "none", "null", "unknown", "-", "--", "tbd", "?"}

# Every written cell is shaded by how sure Netrune is of it, on five steps:
# certain (a serial/MAC match), strong (two keys agree, or a new row), fair (one
# key filling a blank), weak (one key overwriting), doubtful (a forced tiebreak).
# The fills were searched for, not picked: every pair stays at least 15 CIEDE2000
# apart under normal, deutan, protan, and tritan vision (the old three-step
# palette was 6.4 under protan), black text stays at 6.5:1 or better, and the
# warm three darken as doubt grows, so the order survives a greyscale print.
CONFIDENCE_LEVELS = ("doubtful", "weak", "fair", "strong", "certain")
CONFIDENCE_FILLS = {"certain": "AAFCCE", "strong": "7A99E7", "fair": "FDE543", "weak": "D7A569", "doubtful": "ED5C8B"}
CONFIDENCE_SHADES = {"certain": "green", "strong": "blue", "fair": "yellow", "weak": "tan", "doubtful": "red"}
# Weak and doubtful cells are the ones a remark or note marks "(verify)".
VERIFY_LEVELS = {"weak", "doubtful"}
HARDWARE_KEYS = {"serial", "mac"}
STALE_AFTER_DAYS = 30
IPV4 = re.compile(r"(?<![\d.])(\d{1,3}(?:\.\d{1,3}){3})(?![\d.])")


class ExcelError(Exception):
    """A workbook problem the operator can act on."""


def require_openpyxl() -> Any:
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise ExcelError(
            "Excel support needs openpyxl, which is not installed for this Python. "
            "On Arch: sudo pacman -S python-openpyxl  (or: pip install 'netrune[excel]')"
        ) from exc
    return openpyxl


# --------------------------------------------------------------------------
# Value normalization


def _blank(value: Any) -> bool:
    return value is None or str(value).strip().lower() in PLACEHOLDERS


def header_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def header_field(value: Any) -> str:
    key = header_key(value)
    if not key:
        return ""
    for name, synonyms in HEADER_SYNONYMS.items():
        if key in synonyms:
            return name
    # Looser forms such as "Switch Base MAC (Netrune)" or "Mgmt IPv4 addr".
    if "mac" in key and "machine" not in key:
        return "mac"
    if "serial" in key:
        return "serial"
    if "model" in key:
        return "model"
    if "hostname" in key or key.endswith("devicename") or key.endswith("switchname"):
        return "hostname"
    if key.endswith("ip") or "ipaddress" in key or "ipv4" in key:
        return "ip"
    return ""


from netops_platform.identity import name_key  # noqa: E402  (one matching key, ADR 0008)


def ip_key(value: Any) -> str:
    match = IPV4.search(str(value or ""))
    if not match:
        return ""
    try:
        return str(ipaddress.IPv4Address(match.group(1)))
    except ValueError:
        return ""


def serial_key(value: Any) -> str:
    text = re.sub(r"[^A-Z0-9-]", "", str(value or "").upper())
    return "" if text.lower() in PLACEHOLDERS or text in {"NA", "NOTAVAIL", "NOTAVAILABLE"} else text


def identity_key(field_name: str, value: Any) -> str:
    if _blank(value):
        return ""
    if field_name == "hostname":
        return name_key(value)
    if field_name == "ip":
        return ip_key(value)
    if field_name == "mac":
        return normalize_mac(value)
    if field_name == "serial":
        return serial_key(value)
    return str(value).strip().upper()


def mac_style(samples: Iterable[Any]) -> Tuple[str, bool]:
    """The MAC notation a column already uses, so new values match it."""
    counts: Dict[str, int] = {}
    upper = lower = 0
    for sample in samples:
        text = str(sample or "").strip()
        if not normalize_mac(text):
            continue
        if re.fullmatch(r"(?:[0-9A-Fa-f]{2}:){5}[0-9A-Fa-f]{2}", text):
            style = "colon"
        elif re.fullmatch(r"(?:[0-9A-Fa-f]{2}-){5}[0-9A-Fa-f]{2}", text):
            style = "hyphen"
        elif re.fullmatch(r"[0-9A-Fa-f]{4}\.[0-9A-Fa-f]{4}\.[0-9A-Fa-f]{4}", text):
            style = "dotted"
        elif re.fullmatch(r"[0-9A-Fa-f]{6}-[0-9A-Fa-f]{6}", text):
            style = "hyphen-six"
        else:
            style = "bare"
        counts[style] = counts.get(style, 0) + 1
        letters = re.sub(r"[^A-Fa-f]", "", text)
        upper += letters.isupper() and bool(letters)
        lower += letters.islower() and bool(letters)
    style = max(counts, key=counts.get) if counts else "colon"
    return style, upper > lower


def format_mac_as(value: Any, style: str, upper: bool) -> str:
    clean = normalize_mac(value)
    if not clean:
        return ""
    if style == "dotted":
        text = ".".join(clean[i:i + 4] for i in range(0, 12, 4))
    elif style == "hyphen-six":
        text = f"{clean[:6]}-{clean[6:]}"
    elif style == "hyphen":
        text = "-".join(clean[i:i + 2] for i in range(0, 12, 2))
    elif style == "bare":
        text = clean
    else:
        text = ":".join(clean[i:i + 2] for i in range(0, 12, 2))
    return text.upper() if upper else text


# --------------------------------------------------------------------------
# Netrune side


@dataclass
class DeviceIndex:
    devices: List[Dict[str, Any]]
    keys: Dict[str, Dict[str, Set[int]]] = field(default_factory=dict)

    @classmethod
    def build(cls, devices: Sequence[Mapping[str, Any]]) -> "DeviceIndex":
        index = cls([dict(device) for device in devices], {name: {} for name in IDENTITY_FIELDS})
        for position, device in enumerate(index.devices):
            values = {
                "ip": [device.get("ip"), *(device.get("addresses") or [])],
                "mac": [device.get("mac"), *(device.get("macs") or [])],
                "serial": [device.get("serial"), *(device.get("serials") or [])],
                "hostname": [device.get("hostname"), device.get("name"), *(device.get("aliases") or [])],
            }
            for kind, items in values.items():
                for item in items:
                    key = identity_key(kind, item)
                    if key:
                        index.keys[kind].setdefault(key, set()).add(position)
        return index

    def lookup(self, kind: str, value: Any) -> Set[int]:
        key = identity_key(kind, value)
        return set(self.keys.get(kind, {}).get(key, set())) if key else set()


def device_value(device: Mapping[str, Any], field_name: str) -> Any:
    if field_name == "last_seen":
        try:
            seconds = int(device.get("last_seen") or 0)
        except (TypeError, ValueError):
            return None
        # Naive local time: Excel has no time zones, and an operator reads the
        # sheet in the time the box reports.
        return dt.datetime.fromtimestamp(seconds).replace(microsecond=0) if seconds > 0 else None
    value = device.get(field_name)
    if field_name == "hostname":
        value = device.get("hostname") or device.get("name")
    if _blank(value):
        return None
    if field_name == "mac":
        return normalize_mac(value) or None
    if field_name == "model" and not _meaningful_model(value):
        return None
    return str(value).strip()


# --------------------------------------------------------------------------
# Workbook side


@dataclass
class SheetLayout:
    title: str
    header_row: int
    columns: Dict[str, int]
    headers: Dict[int, str]
    first_row: int
    last_row: int
    hostname_fqdn: bool = False
    mac_format: Tuple[str, bool] = ("colon", False)

    def as_dict(self) -> Dict[str, Any]:
        from openpyxl.utils import get_column_letter
        return {
            "sheet": self.title,
            "header_row": self.header_row,
            "data_rows": max(0, self.last_row - self.first_row + 1),
            "columns": {name: f"{get_column_letter(col)} ({self.headers.get(col, '')})" for name, col in sorted(self.columns.items(), key=lambda item: item[1])},
        }


def parse_column_map(spec: str) -> Dict[str, str]:
    """`--map "Mgmt Addr=ip,F=mac"`: header text or column letter to field."""
    mapping: Dict[str, str] = {}
    for part in [item for item in (spec or "").split(",") if item.strip()]:
        if "=" not in part:
            raise ExcelError(f"--map entry '{part}' must look like 'Header=field' or 'C=field'.")
        left, _, right = part.rpartition("=")
        target = right.strip().lower().replace("-", "_")
        if target not in FIELDS + (REMARKS,):
            raise ExcelError(f"--map field '{right.strip()}' is not one of: {', '.join(FIELDS + (REMARKS,))}.")
        mapping[left.strip()] = target
    return mapping


def select_sheets(sheetnames: Sequence[str], spec: str) -> List[str]:
    """`--sheet`: a name (any case/spacing), a 1-based position, or a comma list of either."""
    def resolve(item: str) -> str:
        item = item.strip()
        if item in sheetnames:
            return item
        folded = [name for name in sheetnames if header_key(name) == header_key(item)]
        if len(folded) == 1:
            return folded[0]
        if item.isdigit() and 1 <= int(item) <= len(sheetnames):
            return sheetnames[int(item) - 1]
        listing = ", ".join(f"{number}={name}" for number, name in enumerate(sheetnames, start=1))
        raise ExcelError(f"No sheet named '{item}'. Sheets: {listing}")
    spec = (spec or "").strip()
    if spec in sheetnames:
        return [spec]  # a sheet whose own name has a comma in it
    chosen: List[str] = []
    for item in [part for part in spec.split(",") if part.strip()]:
        name = resolve(item)
        if name not in chosen:
            chosen.append(name)
    if not chosen:
        raise ExcelError("--sheet needs a sheet name or number.")
    return chosen


def detect_layout(ws: Any, overrides: Optional[Mapping[str, str]] = None, scan_rows: int = 30) -> Optional[SheetLayout]:
    """Find the header row: the first of the top rows naming the most fields."""
    from openpyxl.utils import column_index_from_string
    overrides = dict(overrides or {})
    best: Optional[Tuple[int, Dict[str, int], Dict[int, str]]] = None
    max_column = ws.max_column or 0
    for row in range(1, min(scan_rows, ws.max_row or 0) + 1):
        columns: Dict[str, int] = {}
        headers: Dict[int, str] = {}
        for col in range(1, max_column + 1):
            value = ws.cell(row=row, column=col).value
            if value is None or not isinstance(value, str):
                continue
            headers[col] = value.strip()
            name = header_field(value)
            if name and name not in columns:
                columns[name] = col
        for key, target in overrides.items():
            col = 0
            if re.fullmatch(r"[A-Za-z]{1,3}", key):
                col = column_index_from_string(key.upper())
            elif key.isdigit():
                col = int(key)
            else:
                col = next((c for c, text in headers.items() if header_key(text) == header_key(key)), 0)
            if col:
                columns = {name: c for name, c in columns.items() if c != col}
                columns[target] = col
        if not any(name in columns for name in IDENTITY_FIELDS) or len(columns) < 2:
            continue
        if best is None or len(columns) > len(best[1]):
            best = (row, columns, headers)
    if best is None:
        return None
    header_row, columns, headers = best
    last_row = header_row
    for row in range(header_row + 1, (ws.max_row or header_row) + 1):
        if any(not _blank(ws.cell(row=row, column=col).value) for col in columns.values()):
            last_row = row
    layout = SheetLayout(ws.title, header_row, columns, headers, header_row + 1, last_row)
    if "hostname" in columns:
        names = [ws.cell(row=r, column=columns["hostname"]).value for r in range(layout.first_row, last_row + 1)]
        names = [str(n).strip() for n in names if not _blank(n)]
        dotted = sum(1 for n in names if "." in n and not ip_key(n))
        layout.hostname_fqdn = bool(names) and dotted * 2 > len(names)
    if "mac" in columns:
        layout.mac_format = mac_style(ws.cell(row=r, column=columns["mac"]).value for r in range(layout.first_row, last_row + 1))
    return layout


def workbook_features(path: Path) -> List[str]:
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
    except zipfile.BadZipFile as exc:
        raise ExcelError(f"{path} is not an .xlsx/.xlsm workbook (not a zip container).") from exc
    found = []
    for prefix, label in FEATURES_LOST_ON_SAVE.items():
        if any(name.startswith(prefix) for name in names) and label not in found:
            found.append(label)
    return found


def excel_lock_file(path: Path) -> Optional[Path]:
    """Excel's `~$name.xlsx` owner file: the workbook is open somewhere."""
    for candidate in (path.with_name("~$" + path.name), path.with_name("~$" + path.name[2:])):
        if candidate.exists() and candidate != path:
            return candidate
    return None


def load_workbook(path: Path) -> Any:
    openpyxl = require_openpyxl()
    suffix = path.suffix.lower()
    if suffix not in {".xlsx", ".xlsm"}:
        raise ExcelError(f"{path.name}: only .xlsx and .xlsm workbooks are supported"
                         + (" (save .xls files as .xlsx first)." if suffix == ".xls" else "."))
    if not path.is_file():
        raise ExcelError(f"{path} does not exist.")
    try:
        return openpyxl.load_workbook(path, keep_vba=suffix == ".xlsm", rich_text=True)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise ExcelError(f"Could not read {path.name}: {exc}") from exc


# --------------------------------------------------------------------------
# Planning


@dataclass
class CellChange:
    sheet: str
    row: int
    column: int
    header: str
    field: str
    old: Any
    new: Any
    matched_by: str
    device: str
    kind: str = "update"  # update | fill | add | refresh (same value, rewritten)
    confidence: str = "certain"
    why: str = ""

    def as_dict(self) -> Dict[str, Any]:
        from openpyxl.utils import get_column_letter
        return {
            "sheet": self.sheet, "cell": f"{get_column_letter(self.column)}{self.row}", "row": self.row,
            "column": self.header, "field": self.field, "old": _display(self.old), "new": _display(self.new),
            "matched_by": self.matched_by, "device": self.device, "kind": self.kind,
            "confidence": self.confidence, "why": self.why,
        }


@dataclass
class SheetPlan:
    layout: SheetLayout
    changes: List[CellChange] = field(default_factory=list)
    matched_rows: int = 0
    unmatched: List[Dict[str, Any]] = field(default_factory=list)
    conflicts: List[Dict[str, Any]] = field(default_factory=list)
    kept: List[Dict[str, Any]] = field(default_factory=list)
    duplicates: List[Dict[str, Any]] = field(default_factory=list)
    matched_devices: Set[int] = field(default_factory=set)
    additions: List[Dict[str, Any]] = field(default_factory=list)
    ip_checks: List[Dict[str, Any]] = field(default_factory=list)
    remarks_column: int = 0  # set when the sheet had no Remarks column and one is added
    probe_note: str = ""  # what the free-IP probe could not use on this machine


@dataclass
class IpFinding:
    """What a free-IP check concluded about one sheet row's address.

    state: free | in_use | recent (used within STALE_AFTER_DAYS) |
    unconfirmed (names something unknown and silent; awaiting the operator) |
    allocated (the operator said it is still spoken for) | reserved (a
    remark reserves it). `priority` orders allocation suggestions: 1 for an
    empty row, 2 for one last used long ago, 3 for a confirmed old entry.
    """
    state: str
    remark: str
    confidence: str
    why: str
    priority: int = 0
    # What Netrune holds for a device on an empty row's address (name, mac,
    # last_seen): the device is real but missing from the sheet.
    fill: Dict[str, Any] = field(default_factory=dict)


def ip_candidates(ws: Any, layout: SheetLayout, index: Optional["DeviceIndex"] = None) -> List[Dict[str, Any]]:
    """Rows with an address that no Netrune switch matches, by any key.

    `empty` rows name nothing (no hostname, serial, or MAC); the others name
    something Netrune does not know. A row a switch matches is the switch's,
    and plan_sheet fills it instead.
    """
    if "ip" not in layout.columns:
        return []
    found: List[Dict[str, Any]] = []
    for row in range(layout.first_row, layout.last_row + 1):
        cells = {name: ws.cell(row=row, column=col).value for name, col in layout.columns.items()}
        ip = ip_key(cells.get("ip"))
        if not ip or ip.endswith((".0", ".255")):
            continue
        if index is not None and any(index.lookup(kind, cells.get(kind)) for kind in IDENTITY_FIELDS if not _blank(cells.get(kind))):
            continue
        named = [str(cells[name]).strip() for name in ("hostname", "serial", "mac") if not _blank(cells.get(name))]
        found.append({"sheet": layout.title, "row": row, "ip": ip, "empty": not named,
                      "label": named[0] if named else "", "remark": cells.get(REMARKS)})
    return found


def prior_decision(remark: Any) -> str:
    """An earlier answer to "is this old entry free?", kept in the remark itself."""
    parts = {part.lower() for part in _remark_parts(remark)}
    if FREE_OLD_ENTRY.lower() in parts:
        return "accepted"
    if ALLOCATED_OFFLINE.lower() in parts:
        return "rejected"
    return ""


def classify_ip(ip: str, *, empty: bool, known: Mapping[str, str], history: Mapping[str, Tuple[str, int]],
                reply: str, now: float, decision: str = "", details: Optional[Mapping[str, Mapping[str, Any]]] = None) -> IpFinding:
    """Whether one row's address is free to use.

    Never free when a device Netrune knows (online or offline) holds it, or
    one was seen on it within STALE_AFTER_DAYS, or anything answers on it
    (`reply`; every candidate is probed, empty rows included). Otherwise an
    empty row is free, while a row naming something Netrune never saw is free
    only if the operator agrees (`decision` "accepted"); "rejected" keeps it
    allocated, offline. A known device on an empty row carries `details` into
    the finding so its name and MAC can be written into the sheet.
    """
    if known.get(ip):
        return IpFinding("in_use", f"In use ({known[ip]})", "certain", "in inventory",
                         fill=dict((details or {}).get(ip) or {}) if empty else {})
    if reply:
        return IpFinding("in_use", f"In use ({reply})", "certain", reply + (", not in sheet" if empty else ""))
    name, seen = history.get(ip, ("", 0))
    days = int((now - seen) // 86400) if seen else -1
    until = dt.datetime.fromtimestamp(seen).strftime("%Y-%m-%d") if seen else ""
    if seen and days <= STALE_AFTER_DAYS:
        return IpFinding("recent", f"Used by {name} until {until}", "weak", f"used {days} days ago")
    if empty:
        if seen:
            return IpFinding("free", f"Free to use (was {name} until {until})", "fair", f"no reply; last used {days} days ago", 2)
        return IpFinding("free", "Free to use", "certain", "empty row, no reply, no record", 1)
    if decision == "accepted":
        return IpFinding("free", FREE_OLD_ENTRY, "fair", "old entry, confirmed free", 3)
    if decision == "rejected":
        return IpFinding("allocated", ALLOCATED_OFFLINE, "fair", "kept allocated")
    return IpFinding("unconfirmed", "", "weak", "unknown entry, no reply")


def _remark_parts(value: Any) -> List[str]:
    if _blank(value):
        return []
    return [part.strip() for part in re.split(r";\s*", str(value).strip()) if part.strip()]


def _join_labels(fields: Sequence[str]) -> str:
    labels = [REMARK_LABELS.get(name, name) for name in fields]
    return labels[0] if len(labels) == 1 else ", ".join(labels[:-1]) + " and " + labels[-1]


def change_remark(changes: Sequence["CellChange"]) -> str:
    """What one row's writes amount to, as the tech would note it."""
    phrases = []
    renamed = [change for change in changes if change.field == "hostname" and change.kind == "update"]
    if renamed:
        phrases.append(f"Renamed from {_display(renamed[0].old)}")
    updated = [change.field for change in changes if change.kind == "update" and change.field != "hostname"]
    if updated:
        phrases.append("Updated " + _join_labels(updated))
    filled = [change.field for change in changes if change.kind == "fill"]
    if filled:
        phrases.append("Added " + _join_labels(filled))
    text = "; ".join(phrases)
    if text and any(change.confidence in VERIFY_LEVELS for change in changes):
        text += " (verify)"
    return text


def state_remark(device: Mapping[str, Any], now: float) -> str:
    try:
        seen = int(device.get("last_seen") or 0)
    except (TypeError, ValueError):
        seen = 0
    if seen and (now - seen) // 86400 > STALE_AFTER_DAYS:
        return "Last seen " + dt.datetime.fromtimestamp(seen).strftime("%Y-%m-%d")
    return "Offline" if str(device.get("status") or "").upper() == "OFFLINE" else ""


def merge_remark(existing: Any, new: Sequence[str]) -> str:
    """The operator's text kept, this module's own state remarks replaced, new notes appended once."""
    wanted = {phrase.lower() for phrase in new if phrase}
    # Own remarks that still hold keep their place, so a re-run changes nothing.
    merged = [part for part in _remark_parts(existing) if not OWN_REMARK.match(part) or part.lower() in wanted]
    for phrase in new:
        if phrase and phrase.lower() not in {part.lower() for part in merged}:
            merged.append(phrase)
    return "; ".join(merged)


def _display(value: Any) -> str:
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    return "" if value is None else str(value)


def _equivalent(field_name: str, current: Any, desired: Any) -> bool:
    if _blank(current):
        return False
    if field_name in IDENTITY_FIELDS:
        return identity_key(field_name, current) == identity_key(field_name, desired)
    if field_name == "last_seen":
        if isinstance(current, dt.datetime) and isinstance(desired, dt.datetime):
            return abs((current - desired).total_seconds()) < 60
        return False
    have, want = str(current).strip().lower(), str(desired).strip().lower()
    if field_name in {"platform", "manufacturer"}:
        # "Cisco" in a Platform column already agrees with "Cisco IOS", and
        # "Cisco Systems" in a Vendor column with "Cisco".
        return have in want or want in have
    return have == want


def change_confidence(*, matched_by: str, forced: bool, kind: str, field_name: str,
                      device: Mapping[str, Any], now: float) -> Tuple[str, str]:
    """How sure Netrune is of one written cell, and the short reason why.

    The match decides most of it: a serial or base MAC pins the row to one
    switch (certain); two other keys agreeing nearly does (strong); an IP or
    hostname alone can belong to a replaced switch, so it fills a blank at fair
    and overwrites at weak. A --prefer tiebreak is doubtful. A switch Netrune
    has not heard from lately drops any of it one step.
    """
    keys = set(matched_by.split("+"))
    hardware = sorted(keys & HARDWARE_KEYS, key=lambda key: key != "serial")
    match = f"{hardware[0]} match" if hardware else f"{len(keys)} keys match" if len(keys) > 1 else f"{matched_by} only"
    if kind == "refresh":
        return "certain", "same value, reformatted"
    if kind == "add":
        level, why = 3, "new row from inventory"
    elif forced:
        level, why = 0, "--prefer overrode a conflict"
    elif kind == "update" and field_name in HARDWARE_KEYS and (hardware or len(keys) > 1):
        # A sheet serial or MAC that disagrees with a switch pinned by its
        # other keys is a typo or a hardware swap: either way, glance at it.
        level, why = 2, f"{match}; replaced value"
    elif hardware or len(keys) > 1:
        level, why = (4 if hardware else 3), f"{match}; " + ("blank filled" if kind == "fill" else "replaced value")
    elif kind == "fill":
        level, why = 2, f"{match}; blank filled"
    else:
        level, why = 1, f"{match}; replaced value"
    try:
        seen = int(device.get("last_seen") or 0)
    except (TypeError, ValueError):
        seen = 0
    age_days = int((now - seen) // 86400) if seen > 0 else -1
    if age_days < 0 or age_days > STALE_AFTER_DAYS:
        level = max(0, level - 1)
        why = "never polled" if age_days < 0 else f"last seen {age_days} days ago"
    return CONFIDENCE_LEVELS[level], why


def _cell_text_for(layout: SheetLayout, field_name: str, value: Any) -> Any:
    if field_name == "mac":
        style, upper = layout.mac_format
        return format_mac_as(value, style, upper)
    if field_name == "hostname" and not layout.hostname_fqdn:
        return _host_label(value)
    return value


def plan_sheet(
    ws: Any,
    layout: SheetLayout,
    index: DeviceIndex,
    *,
    fields: Sequence[str] = DEFAULT_FIELDS,
    fill_only: bool = False,
    prefer: str = "",
    add_missing: bool = False,
    refresh: Sequence[str] = (),
    now: Optional[float] = None,
    remarks: bool = False,
    ip_findings: Optional[Mapping[int, IpFinding]] = None,
    row_prefer: Optional[Mapping[int, Mapping[str, str]]] = None,
) -> SheetPlan:
    """Plan the cell writes for one sheet.

    `refresh` names columns to rewrite on every matched row, even where the
    cell already holds the same value in another notation, and even under
    `fill_only`. A column being refreshed is one the operator does not trust,
    so it only matches rows that have no other key to go by.

    `remarks` adds one brief note per row to the Remarks column (adding the
    column when the sheet has none): what was changed, a switch that is
    offline or long unseen, a row that conflicts or repeats another, and, from
    `ip_findings` (keyed by row), whether the row's address is free to use.
    """
    from openpyxl.cell.cell import MergedCell
    plan = SheetPlan(layout)
    refresh = set(refresh)
    now = time.time() if now is None else now
    writable = [name for name in FIELDS if name in layout.columns and (name in fields or name in refresh)]
    owners: Dict[int, List[int]] = {}
    row_device: Dict[int, Dict[str, Any]] = {}
    row_notes: Dict[int, List[Tuple[str, str, str]]] = {}  # row -> (remark, confidence, why)
    for row in range(layout.first_row, layout.last_row + 1):
        cells = {name: ws.cell(row=row, column=col).value for name, col in layout.columns.items()}
        if all(_blank(cells.get(name)) for name in IDENTITY_FIELDS):
            continue
        votes: Dict[str, int] = {}
        ambiguous: List[str] = []
        for kind in IDENTITY_FIELDS:
            if kind not in layout.columns or _blank(cells.get(kind)):
                continue
            hits = index.lookup(kind, cells[kind])
            if len(hits) == 1:
                votes[kind] = next(iter(hits))
            elif len(hits) > 1:
                ambiguous.append(kind)
        label = next((str(cells[k]) for k in ("hostname", "ip", "serial", "mac") if not _blank(cells.get(k))), f"row {row}")
        trusted = {kind: pos for kind, pos in votes.items() if kind not in refresh}
        if trusted:
            votes = trusted
        distinct = set(votes.values())
        forced = False
        if not distinct:
            plan.unmatched.append({"sheet": layout.title, "row": row, "row_label": label,
                                   "reason": "ambiguous " + "/".join(ambiguous) if ambiguous else "no Netrune device has these values"})
            finding = (ip_findings or {}).get(row)
            if finding:
                plan.ip_checks.append({"sheet": layout.title, "row": row, "ip": ip_key(cells.get("ip")),
                                       "named": label if label != ip_key(cells.get("ip")) else "",
                                       "state": finding.state, "remark": finding.remark, "why": finding.why,
                                       "priority": finding.priority})
                for name in ("hostname", "mac"):
                    value = finding.fill.get(name)
                    column = layout.columns.get(name)
                    if not value or not column or not _blank(ws.cell(row=row, column=column).value):
                        continue
                    level, why = change_confidence(matched_by="ip", forced=False, kind="fill", field_name=name,
                                                   device=finding.fill, now=now)
                    plan.changes.append(CellChange(layout.title, row, column, layout.headers.get(column, name), name, None,
                                                   _cell_text_for(layout, name, value), "ip", finding.remark[8:-1],
                                                   "fill", level, why))
                if finding.state == "unconfirmed":
                    # Nothing new is said until the operator answers, but a
                    # "Free to use" from when the row was empty is cleared.
                    row_notes.setdefault(row, []).append(("", "certain", "cleared old remark"))
                    continue
                row_notes.setdefault(row, []).append((finding.remark, finding.confidence, finding.why))
            elif ip_findings is not None:
                # Not an open address any more (someone named a device on it):
                # an old "Free to use" here is cleared.
                row_notes.setdefault(row, []).append(("", "certain", "cleared old remark"))
            continue
        if len(distinct) > 1:
            # An answer given for this row (`health review`) counts only while
            # the row still reads the same, so a stale one never lands elsewhere.
            answered = (row_prefer or {}).get(row) or {}
            chosen = answered.get("kind") if answered.get("label") == label else ""
            chosen = chosen if chosen in votes else prefer
            if chosen and chosen in votes:
                distinct = {votes[chosen]}
                forced = True
            else:
                plan.conflicts.append({
                    "sheet": layout.title, "row": row, "row_label": label,
                    "points_to": ", ".join(f"{kind}->{_device_label(index.devices[pos])}" for kind, pos in votes.items()),
                    "votes": {kind: _device_label(index.devices[pos]) for kind, pos in votes.items()},
                })
                keys_by_device: Dict[int, List[str]] = {}
                for kind, pos in votes.items():
                    keys_by_device.setdefault(pos, []).append(REMARK_LABELS[kind])
                said = ", ".join(f"{'/'.join(keys)} is {_host_label(index.devices[pos].get('hostname') or index.devices[pos].get('name') or index.devices[pos].get('ip'))}"
                                 for pos, keys in keys_by_device.items())
                row_notes.setdefault(row, []).append((f"Check: {said}", "doubtful", "keys disagree"))
                continue
        position = next(iter(distinct))
        device = index.devices[position]
        matched_by = "+".join(kind for kind, pos in votes.items() if pos == position)
        plan.matched_rows += 1
        plan.matched_devices.add(position)
        owners.setdefault(position, []).append(row)
        row_device[row] = device
        for name in writable:
            desired = device_value(device, name)
            if desired is None:
                if name in refresh and not _blank(ws.cell(row=row, column=layout.columns[name]).value):
                    plan.kept.append({"sheet": layout.title, "row": row, "column": layout.headers.get(layout.columns[name], name),
                                      "reason": f"--refresh: Netrune has no {name} for it"})
                continue
            column = layout.columns[name]
            cell = ws.cell(row=row, column=column)
            if isinstance(cell, MergedCell):
                continue
            current = cell.value
            if cell.data_type == "f" or (isinstance(current, str) and current.startswith("=")):
                if not _equivalent(name, current, desired):
                    plan.kept.append({"sheet": layout.title, "row": row, "column": layout.headers.get(column, name), "reason": "formula cell left as-is"})
                continue
            text = _cell_text_for(layout, name, desired)
            if _equivalent(name, current, desired):
                if name not in refresh or _display(current) == _display(text):
                    continue
                kind = "refresh"
            elif fill_only and not _blank(current) and name not in refresh:
                plan.kept.append({"sheet": layout.title, "row": row, "column": layout.headers.get(column, name),
                                  "reason": f"--fill-only kept '{_display(current)}' (Netrune: '{_display(text)}')"})
                continue
            else:
                kind = "fill" if _blank(current) else "update"
            confidence, why = change_confidence(matched_by=matched_by, forced=forced, kind=kind,
                                                field_name=name, device=device, now=now)
            plan.changes.append(CellChange(
                layout.title, row, column, layout.headers.get(column, name), name, current,
                text, matched_by, _device_label(device), kind, confidence, why,
            ))
    for position, rows in owners.items():
        if len(rows) > 1:
            plan.duplicates.append({"sheet": layout.title, "device": _device_label(index.devices[position]), "rows": ", ".join(map(str, rows))})
            for row in rows[1:]:
                row_notes.setdefault(row, []).append((f"Duplicate of row {rows[0]}", "fair", "switch on several rows"))
    if add_missing:
        next_row = layout.last_row + 1
        for position, device in enumerate(index.devices):
            if position in plan.matched_devices:
                continue
            values = {name: device_value(device, name) for name in writable}
            if not any(values.get(name) is not None for name in IDENTITY_FIELDS):
                continue
            plan.additions.append({"row": next_row, "device": _device_label(device)})
            for name in writable:
                if values.get(name) is None:
                    continue
                column = layout.columns[name]
                confidence, why = change_confidence(matched_by="", forced=False, kind="add",
                                                    field_name=name, device=device, now=now)
                plan.changes.append(CellChange(
                    layout.title, next_row, column, layout.headers.get(column, name), name, None,
                    _cell_text_for(layout, name, values[name]), "not in sheet", _device_label(device), "add",
                    confidence, why,
                ))
            next_row += 1
    if remarks:
        _plan_remarks(ws, plan, row_device, row_notes, now)
    return plan


def _plan_remarks(ws: Any, plan: SheetPlan, row_device: Mapping[int, Mapping[str, Any]],
                  row_notes: Dict[int, List[Tuple[str, str, str]]], now: float) -> None:
    from openpyxl.cell.cell import MergedCell
    layout = plan.layout
    by_row: Dict[int, List[CellChange]] = {}
    for change in plan.changes:
        by_row.setdefault(change.row, []).append(change)
    for row, changes in by_row.items():
        noted = change_remark([change for change in changes if change.kind != "refresh"])
        if any(change.kind == "add" for change in changes):
            row_notes.setdefault(row, []).append(("Added, found on network", "strong", "new row from inventory"))
        elif noted:
            worst = min((change.confidence for change in changes), key=CONFIDENCE_LEVELS.index)
            row_notes.setdefault(row, []).append((noted, worst, "row updated"))
    for row, device in row_device.items():
        status = state_remark(device, now)
        # Run for every matched row, so a stale "Offline" or "Free to use" is cleared.
        row_notes.setdefault(row, []).append((status, "fair" if status else "certain", "switch state"))
    column = layout.columns.get(REMARKS)
    if not column:
        if not any(note for notes in row_notes.values() for note, _, _ in notes):
            return
        column = max(layout.headers) + 1 if layout.headers else 1
        while any(not _blank(ws.cell(row=r, column=column).value) for r in range(layout.header_row, layout.last_row + 1)):
            column += 1
        plan.remarks_column = column
        plan.changes.append(CellChange(layout.title, layout.header_row, column, REMARKS_HEADER, REMARKS, None,
                                       REMARKS_HEADER, "", "", "header", "certain", "new Remarks column"))
    for row in sorted(row_notes):
        cell = ws.cell(row=row, column=column)
        if isinstance(cell, MergedCell) or cell.data_type == "f":
            continue
        notes = [(text, level, why) for text, level, why in row_notes[row]]
        finding_free = any(text.lower().startswith("free to use") for text, _, _ in notes)
        theirs = [part for part in _remark_parts(cell.value) if not OWN_REMARK.match(part)]
        if finding_free and any(RESERVED_WORDS.search(part) for part in theirs):
            plan.kept.append({"sheet": layout.title, "row": row, "column": layout.headers.get(column, REMARKS_HEADER),
                              "reason": f"no reply, but the remark says '{'; '.join(theirs)[:40]}'"})
            notes = [note for note in notes if not note[0].lower().startswith("free to use")]
            for check in plan.ip_checks:
                if check["row"] == row:
                    check.update(state="reserved", remark="(kept: " + "; ".join(theirs)[:40] + ")", why="remark reserves it")
        merged = merge_remark(cell.value, [text for text, _, _ in notes])
        if merged == ("; ".join(_remark_parts(cell.value))) or (not merged and _blank(cell.value)):
            continue
        written = [(text, level, why) for text, level, why in notes if text] or [("", "certain", "cleared old remark")]
        level, why = min(((lvl, w) for _, lvl, w in written), key=lambda item: CONFIDENCE_LEVELS.index(item[0]))
        plan.changes.append(CellChange(layout.title, row, column, layout.headers.get(column, REMARKS_HEADER), REMARKS,
                                       cell.value, merged or None, "", "", "fill" if _blank(cell.value) else "update",
                                       level, why))


def _device_label(device: Mapping[str, Any]) -> str:
    name = _host_label(device.get("hostname") or device.get("name") or "")
    return f"{name} ({device.get('ip')})" if name else str(device.get("ip") or "?")


# --------------------------------------------------------------------------
# Applying


def apply_plans(wb: Any, plans: Sequence[SheetPlan], *, shade: bool = True, comments: bool = False,
                stamp: str = "") -> int:
    from openpyxl.comments import Comment
    from openpyxl.formula.translate import Translator
    from openpyxl.styles import PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import TableColumn
    fills = {level: PatternFill(start_color=color, end_color=color, fill_type="solid") for level, color in CONFIDENCE_FILLS.items()}
    stamp = stamp or dt.datetime.now().strftime("%Y-%m-%d")
    author = _operator_name()
    applied = 0
    for plan in plans:
        ws = wb[plan.layout.title]
        template_row = plan.layout.last_row
        added_rows = sorted({change.row for change in plan.changes if change.kind == "add"})
        for row in added_rows:
            # A new row looks and calculates like the one above it: same
            # styles, and that row's formulas re-pointed at this row.
            for col in range(1, (ws.max_column or 0) + 1):
                source = ws.cell(row=template_row, column=col)
                target = ws.cell(row=row, column=col)
                if source.has_style:
                    target._style = copy.copy(source._style)
                if source.data_type == "f" and isinstance(source.value, str):
                    origin = f"{get_column_letter(col)}{template_row}"
                    target.value = Translator(source.value, origin=origin).translate_formula(f"{get_column_letter(col)}{row}")
        if plan.remarks_column:
            # The new Remarks column looks like the header and the column beside it.
            left = plan.remarks_column - 1
            for row in range(plan.layout.header_row, max([plan.layout.last_row] + added_rows) + 1):
                source = ws.cell(row=row, column=left) if left >= 1 else None
                if source is not None and source.has_style:
                    ws.cell(row=row, column=plan.remarks_column)._style = copy.copy(source._style)
            letter = get_column_letter(plan.remarks_column)
            if left >= 1 and ws.column_dimensions[get_column_letter(left)].width:
                ws.column_dimensions[letter].width = max(ws.column_dimensions[get_column_letter(left)].width, 30)
            else:
                ws.column_dimensions[letter].width = 30
            for table in ws.tables.values():
                start, _, end = table.ref.partition(":")
                end_col = re.sub(r"[0-9]", "", end)
                start_row = int(re.sub(r"[A-Z]", "", start))
                if start_row == plan.layout.header_row and end_col == get_column_letter(left):
                    table.ref = f"{start}:{letter}{re.sub(r'[A-Z]', '', end)}"
                    table.tableColumns.append(TableColumn(id=max((c.id for c in table.tableColumns), default=0) + 1, name=REMARKS_HEADER))
                    if table.autoFilter is not None:
                        table.autoFilter.ref = table.ref
        for change in plan.changes:
            cell = ws.cell(row=change.row, column=change.column)
            cell.value = change.new
            if isinstance(change.new, dt.datetime) and cell.number_format in {"General", ""}:
                cell.number_format = "yyyy-mm-dd hh:mm"
            if change.kind == "header":
                applied += 1
                continue
            if comments and change.field != REMARKS:
                note = ("Added, found on network" if change.kind == "add"
                        else f"Was '{_display(change.old)}'" if change.kind in {"update", "refresh"}
                        else "Filled in")
                cell.comment = Comment(f"{stamp}: {note}" + (" (verify)" if change.confidence in VERIFY_LEVELS else ""),
                                       author, width=220, height=50)
            if shade:
                cell.fill = fills[change.confidence]
            applied += 1
        if added_rows:
            last = max(added_rows)
            for table in ws.tables.values():
                start, _, end = table.ref.partition(":")
                end_row = int(re.sub(r"[A-Z]", "", end))
                start_row = int(re.sub(r"[A-Z]", "", start))
                if start_row <= plan.layout.header_row <= end_row and end_row >= plan.layout.last_row:
                    table.ref = f"{start}:{re.sub(r'[0-9]', '', end)}{last}"
                    if table.autoFilter is not None:
                        table.autoFilter.ref = table.ref
    return applied


def _operator_name() -> str:
    import getpass
    try:
        return getpass.getuser()
    except (KeyError, OSError):  # pragma: no cover - no passwd entry
        return "Network team"


def verify_saved(path: Path, plans: Sequence[SheetPlan]) -> List[str]:
    """Reopen what was written and confirm every planned cell holds its value."""
    wb = load_workbook(path)
    problems = []
    for plan in plans:
        ws = wb[plan.layout.title]
        for change in plan.changes:
            value = ws.cell(row=change.row, column=change.column).value
            if isinstance(change.new, dt.datetime):
                ok = isinstance(value, dt.datetime) and abs((value - change.new).total_seconds()) < 1
            else:
                ok = str(value) == str(change.new)
            if not ok:
                problems.append(f"{change.sheet}!{change.as_dict()['cell']}: expected '{_display(change.new)}', found '{_display(value)}'")
    return problems


# --------------------------------------------------------------------------
# A fresh tracker workbook


TRACKER_STATUSES = ("Not started", "Planned", "In progress", "Blocked", "Done", "Cancelled")
TRACKER_CHANGES = ("Hardware swap", "Firmware upgrade", "VLAN change", "Readdress", "Port change",
                   "Decommission", "New install", "Other")


def build_tracker(devices: Sequence[Mapping[str, Any]], path: Path, *, generated: str = "",
                  migration_rows: int = 300) -> Dict[str, Any]:
    """Asset inventory, live summary, and a migration/change tracker in one workbook.

    Everything after the data is formulas (COUNTIFS, INDEX/MATCH, TODAY), so
    the summary and tracker keep calculating as the operator edits statuses,
    owners, and dates -- and `netrune excel sync` can refresh the Inventory
    sheet later without disturbing them.
    """
    openpyxl = require_openpyxl()
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo

    generated = generated or dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    title_font = Font(bold=True, size=14, color="1F3864")
    muted = Font(italic=True, color="595959")
    red = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
    amber = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
    green = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")
    grey = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    thin = Side(style="thin", color="BFBFBF")

    wb = openpyxl.Workbook()
    inv = wb.active
    inv.title = "Inventory"
    columns = ["Hostname", "IP Address", "Manufacturer", "Platform", "Model", "Serial Number", "Base MAC", "Status",
               "Last Seen", "Days Since Seen", "Site", "Owner", "Notes"]
    inv.append(columns)
    rows = sorted(devices, key=lambda d: (_host_label(d.get("hostname") or d.get("name") or "").lower(), str(d.get("ip") or "")))
    for number, device in enumerate(rows, start=2):
        inv.append([
            _host_label(device.get("hostname") or device.get("name") or "") or None,
            device.get("ip") or None,
            device_value(device, "manufacturer"),
            device_value(device, "platform"),
            device_value(device, "model"),
            device_value(device, "serial"),
            format_mac_as(device.get("mac"), "colon", False) or None,
            device_value(device, "status"),
            device_value(device, "last_seen"),
            f'=IF(I{number}="","",INT(TODAY()-I{number}))',
            None, None, None,
        ])
        inv.cell(row=number, column=9).number_format = "yyyy-mm-dd hh:mm"
    last = max(2, len(rows) + 1)
    table = Table(displayName="Inventory", ref=f"A1:{get_column_letter(len(columns))}{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    inv.add_table(table)
    inv.freeze_panes = "B2"
    for col, width in zip("ABCDEFGHIJKLM", (30, 15, 14, 14, 24, 16, 19, 10, 17, 10, 16, 14, 40)):
        inv.column_dimensions[col].width = width
    inv.conditional_formatting.add(f"H2:H{last}", CellIsRule(operator="equal", formula=['"OFFLINE"'], fill=red))
    inv.conditional_formatting.add(f"H2:H{last}", CellIsRule(operator="equal", formula=['"ONLINE"'], fill=green))
    inv.conditional_formatting.add(f"J2:J{last}", FormulaRule(formula=[f'AND(ISNUMBER(J2),J2>7)'], fill=amber))

    summary = wb.create_sheet("Summary", 0)
    summary["A1"] = "Network inventory summary"
    summary["A1"].font = title_font
    summary["A2"] = f"Generated by Netrune {generated}. Counts are live formulas over the Inventory sheet."
    summary["A2"].font = muted
    rng = lambda col: f"Inventory!${col}$2:${col}${last}"  # noqa: E731
    kpis = [
        ("Devices", f'=COUNTA({rng("A")})'),
        ("Online", f'=COUNTIF({rng("H")},"ONLINE")'),
        ("Offline", f'=COUNTIF({rng("H")},"OFFLINE")'),
        ("Online %", "=IF(B4=0,0,B5/B4)"),
        ("Not seen for 7+ days", f'=COUNTIF({rng("J")},">7")'),
        ("Missing serial", f'=COUNTBLANK({rng("F")})'),
        ("Missing base MAC", f'=COUNTBLANK({rng("G")})'),
    ]
    for offset, (label, formula) in enumerate(kpis, start=4):
        summary.cell(row=offset, column=1, value=label).font = Font(bold=True)
        summary.cell(row=offset, column=2, value=formula)
    summary["B7"].number_format = "0.0%"
    models = sorted({str(device_value(d, "model")) for d in rows if device_value(d, "model")})
    makers = sorted({str(d.get("manufacturer")) for d in rows if not _blank(d.get("manufacturer"))})
    platforms = sorted({str(d.get("platform")) for d in rows if not _blank(d.get("platform"))})
    start = 13
    for col, title in ((1, "Model"), (2, "Devices"), (3, "Online"), (4, "Offline")):
        cell = summary.cell(row=start - 1, column=col, value=title)
        cell.font, cell.fill = header_font, header_fill
    for offset, model in enumerate(models):
        row = start + offset
        summary.cell(row=row, column=1, value=model)
        summary.cell(row=row, column=2, value=f"=COUNTIF({rng('E')},A{row})")
        summary.cell(row=row, column=3, value=f'=COUNTIFS({rng("E")},A{row},{rng("H")},"ONLINE")')
        summary.cell(row=row, column=4, value=f'=COUNTIFS({rng("E")},A{row},{rng("H")},"OFFLINE")')
    if models:
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Devices by model"
        chart.style = 10
        chart.height = max(7.5, 0.5 * len(models) + 3)
        chart.width = 16
        data = Reference(summary, min_col=2, min_row=start - 1, max_row=start + len(models) - 1)
        cats = Reference(summary, min_col=1, min_row=start, max_row=start + len(models) - 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.legend = None
        summary.add_chart(chart, "F4")
    mrow = start + len(models) + 2
    for col, title in ((1, "Manufacturer"), (2, "Devices")):
        cell = summary.cell(row=mrow, column=col, value=title)
        cell.font, cell.fill = header_font, header_fill
    for offset, maker in enumerate(makers, start=1):
        summary.cell(row=mrow + offset, column=1, value=maker)
        summary.cell(row=mrow + offset, column=2, value=f"=COUNTIF({rng('C')},A{mrow + offset})")
    prow = mrow + len(makers) + 2
    for col, title in ((1, "Platform"), (2, "Devices")):
        cell = summary.cell(row=prow, column=col, value=title)
        cell.font, cell.fill = header_font, header_fill
    for offset, platform in enumerate(platforms, start=1):
        summary.cell(row=prow + offset, column=1, value=platform)
        summary.cell(row=prow + offset, column=2, value=f"=COUNTIF({rng('D')},A{prow + offset})")
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 12

    mig = wb.create_sheet("Migration Tracker")
    mig_columns = ["Device", "IP Address", "Current Model", "Change Type", "Target / Details", "Owner",
                   "Planned Date", "Status", "Completed Date", "Days Overdue", "Notes"]
    mig.append(mig_columns)
    end = migration_rows + 1
    for row in range(2, end + 1):
        lookup = f"MATCH(A{row},Inventory!$A$2:$A${last},0)"
        mig.cell(row=row, column=2, value=f'=IF(A{row}="","",IFERROR(INDEX(Inventory!$B$2:$B${last},{lookup}),"not in inventory"))')
        mig.cell(row=row, column=3, value=f'=IF(A{row}="","",IFERROR(INDEX(Inventory!$E$2:$E${last},{lookup}),""))')
        mig.cell(row=row, column=10, value=f'=IF(OR(G{row}="",H{row}="Done",H{row}="Cancelled"),"",MAX(0,INT(TODAY()-G{row})))')
        mig.cell(row=row, column=7).number_format = "yyyy-mm-dd"
        mig.cell(row=row, column=9).number_format = "yyyy-mm-dd"
    mig_table = Table(displayName="MigrationTracker", ref=f"A1:{get_column_letter(len(mig_columns))}{end}")
    mig_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
    mig.add_table(mig_table)
    mig.freeze_panes = "B2"
    for col, width in zip("ABCDEFGHIJK", (30, 15, 22, 18, 36, 14, 13, 13, 15, 12, 40)):
        mig.column_dimensions[col].width = width
    device_list = DataValidation(type="list", formula1=f"=Inventory!$A$2:$A${last}", allow_blank=True,
                                 showErrorMessage=False)
    status_list = DataValidation(type="list", formula1='"' + ",".join(TRACKER_STATUSES) + '"', allow_blank=True)
    change_list = DataValidation(type="list", formula1='"' + ",".join(TRACKER_CHANGES) + '"', allow_blank=True)
    dates = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
    for validation, cells in ((device_list, f"A2:A{end}"), (change_list, f"D2:D{end}"),
                              (status_list, f"H2:H{end}"), (dates, f"G2:G{end}"), (dates, f"I2:I{end}")):
        if validation not in mig.data_validations.dataValidation:
            mig.add_data_validation(validation)
        validation.add(cells)
    mig.conditional_formatting.add(f"A2:K{end}", FormulaRule(formula=['$H2="Done"'], fill=green))
    mig.conditional_formatting.add(f"A2:K{end}", FormulaRule(formula=['$H2="Blocked"'], fill=red))
    mig.conditional_formatting.add(f"A2:K{end}", FormulaRule(formula=['$H2="Cancelled"'], fill=grey))
    mig.conditional_formatting.add(f"J2:J{end}", FormulaRule(formula=['AND(ISNUMBER(J2),J2>0)'], fill=amber))
    mig.conditional_formatting.add(f"B2:B{end}", CellIsRule(operator="equal", formula=['"not in inventory"'], fill=red))

    trow = prow + len(platforms) + 3
    summary.cell(row=trow, column=1, value="Migration progress").font = title_font
    for col, title in ((1, "Status"), (2, "Changes")):
        cell = summary.cell(row=trow + 1, column=col, value=title)
        cell.font, cell.fill = header_font, header_fill
    for offset, status in enumerate(TRACKER_STATUSES, start=2):
        summary.cell(row=trow + offset, column=1, value=status)
        summary.cell(row=trow + offset, column=2, value=f"=COUNTIF('Migration Tracker'!$H$2:$H${end},A{trow + offset})")
    done_row = trow + len(TRACKER_STATUSES) + 2
    summary.cell(row=done_row, column=1, value="Complete %").font = Font(bold=True)
    summary.cell(row=done_row, column=2, value=(
        f"=IFERROR(COUNTIF('Migration Tracker'!$H$2:$H${end},\"Done\")/"
        f"(COUNTA('Migration Tracker'!$H$2:$H${end})-COUNTIF('Migration Tracker'!$H$2:$H${end},\"Cancelled\")),0)"))
    summary.cell(row=done_row, column=2).number_format = "0.0%"
    summary.cell(row=done_row + 1, column=1, value="Overdue").font = Font(bold=True)
    summary.cell(row=done_row + 1, column=2, value=f"=COUNTIF('Migration Tracker'!$J$2:$J${end},\">0\")")
    for row in summary.iter_rows(min_row=start - 1, max_row=start + len(models) - 1, max_col=4):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="center")

    wb.properties.creator = "Netrune"
    wb.properties.title = "Network inventory and migration tracker"
    wb.save(path)
    return {"devices": len(rows), "models": len(models), "migration_rows": migration_rows,
            "sheets": [ws.title for ws in wb.worksheets]}

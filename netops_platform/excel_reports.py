"""Fixed-definition Excel reports built from Netrune's memory.

Every report here has one written definition (REPORTS[name].rule) that says
exactly which records it lists and how it decides; nothing is inferred per
run. The operator specifies new reports; each becomes one entry in REPORTS
and one collector below. The definitions are also written into the workbook's
"Report Info" sheet, so a reader can see how every row got there.

`collect(app, since)` reads the database once into plain rows; `build(data,
path, names)` writes the workbook. Both are free functions; excel_cli.py owns
the command lifecycle.
"""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

from netops_platform import excel
from netops_platform.devices import ap_model, device_manufacturer, is_switch_row
from netops_platform.identity import _host_label, neighbor_port_transitions, normalize_mac
from netops_platform.inventory_ingest import _meaningful_model, aruba_model_name


# Within this many seconds, changes at one position are one event (one scan).
SAME_SCAN_SECONDS = 600
DEFAULT_SINCE_DAYS = 30


@dataclass(frozen=True)
class Report:
    name: str
    sheet: str
    columns: Tuple[str, ...]
    rule: str


REPORTS: Dict[str, Report] = {report.name: report for report in (
    Report("active-aps", "Active APs",
           ("Name", "IP Address", "MAC", "Manufacturer", "Model", "Switch", "Switch Port", "Last Seen"),
           "Every access point whose status is ONLINE at Netrune's last check. Rows the AP inventory picked up "
           "that are really switches (an Aruba 'Switch, revision' banner, or a switch model such as C9200CX) are "
           "left out. Model is the AP's model number when the controller or LLDP reported one, else blank."),
    Report("active-switches", "Active Switches",
           ("Hostname", "IP Address", "Manufacturer", "Model", "Serial Number", "Base MAC", "Platform", "Last Seen"),
           "Every switch in the inventory whose status is ONLINE at Netrune's last check."),
    Report("transitions", "Transitions",
           ("When", "Type", "Change", "Device", "Kept The Same", "From", "To", "Switch / Port"),
           "Changes in the period, one row per change: Renamed (same MAC or IP, a different host name; adding or "
           "dropping the domain does not count), IP changed (same MAC, new address), IP now on another device "
           "(same address, different MAC), Serial changed, Model changed (a different model, not a relabel), "
           "Moved (an AP seen on a different switch or port), and Replaced at neighbour port / Renamed (same MAC) / "
           "Replaced or renamed (no MAC to compare): a neighbouring switch's port that saw one device now sees "
           "another of the same kind, with how many other links on that switch changed in the same scan (none is "
           "the strongest evidence). Ports shared by several neighbours, ports flapping between two, phones, and "
           "hosts known only by address are left out. Status, platform-label and source changes are "
           "left out, as are values that were blank, N/A, or a MAC standing in for a name. A change Netrune "
           "recorded under several keys is listed once, with every key that stayed the same."),
    Report("switch-history", "Switch History",
           ("Hostname", "IP Address", "State", "Manufacturer", "Model", "Serial Number", "Base MAC",
            "First Seen", "Last Seen", "Hardware Change", "Change Date"),
           "Every switch Netrune holds, online or offline, plus a row for each piece of hardware replaced in the "
           "period. A replacement changes the box itself: at the same hostname or IP, in one scan, the base MAC "
           "changes together with the serial number or the model (or, when no MAC is known, the serial and the "
           "model both change). A serial that changes while the base MAC stays is the same switch read "
           "differently (Cisco reports a motherboard and a system serial), and a serial Netrune has seen on more "
           "than one switch is not evidence of anything. Renames, relabels, and a value appearing where there was "
           "none are not replacements; a swap that reverts later is treated as a reading error and left out. A switch "
           "that replaced another under a new name and address is caught from the neighbour tables: a neighbouring "
           "port that saw the old switch now sees it, the old one is gone, and the rest of that neighbour barely "
           "changed (confidence 75 or more)."),
    Report("ap-history", "AP History",
           ("Name", "IP Address", "State", "Manufacturer", "Model", "MAC", "Switch", "Switch Port",
            "First Seen", "Last Seen", "Hardware Change", "Change Date"),
           "Every AP Netrune holds, online or offline (switch rows left out as in Active APs), plus a row for "
           "each AP replaced in the period. A replacement is a different MAC at the same AP name, or a different "
           "model or maker at that name (Ruckus to Cisco, C9105AXW to C9124AXE); the controller alternating "
           "between a model number and a software banner is not one. IP changes are not replacements: AP "
           "addresses move with DHCP. A swap that reverts later is left out."),
)}


# --------------------------------------------------------------------------
# Value helpers

_MAC_LIKE = re.compile(r"^(?:[0-9a-f]{2}[ :.-]?){5}[0-9a-f]{2}$|^(?:[0-9a-f]{4}\.){2}[0-9a-f]{4}$", re.IGNORECASE)
def placeholder_name(value: Any) -> bool:
    text = str(value or "").strip()
    return excel._blank(text) or bool(_MAC_LIKE.match(text))


def real_serial(value: Any) -> str:
    return excel.serial_key(value) if not excel._blank(value) else ""


def real_mac(value: Any) -> str:
    return normalize_mac(value) if not excel._blank(value) else ""


def switch_model_changed(old: Any, new: Any) -> bool:
    """A different switch model, not the same one spelled another way.

    `JL075A` (an HPE product number alone) and `3810M-16SFP+-2-slot` are one
    model, as are `Aruba JL258A 2930F-8G-PoE+-2SFP+ Switch` and
    `2930F-8G-PoE+-2SFP+`; a different maker always is a different model.
    """
    makers = (device_manufacturer(model=old, platform=old), device_manufacturer(model=new, platform=new))
    if all(makers) and makers[0] != makers[1]:
        return True
    before, after = (re.sub(r"^WS-", "", aruba_model_name(value).upper()) for value in (old, new))
    # "C2960S" is a short label for "WS-C2960S-24TD-L", not another model.
    return bool(before and after and before not in after and after not in before)


def _maker(platform: Any, model: Any = "") -> str:
    return device_manufacturer(platform=platform, model=model) or ""


def _when(epoch: Any) -> Optional[dt.datetime]:
    try:
        seconds = int(epoch or 0)
    except (TypeError, ValueError):
        return None
    return dt.datetime.fromtimestamp(seconds).replace(microsecond=0) if seconds > 0 else None


def _name(value: Any) -> str:
    return _host_label(value) if not placeholder_name(value) else ""


def parse_since(value: str, now: float) -> Tuple[int, str]:
    """`30d`, `12w`, `6m`, `all`, or `YYYY-MM-DD` -> (epoch, label)."""
    text = (value or f"{DEFAULT_SINCE_DAYS}d").strip().lower()
    if text == "all":
        return 0, "all history"
    match = re.fullmatch(r"(\d+)\s*([dwm])", text)
    if match:
        days = int(match.group(1)) * {"d": 1, "w": 7, "m": 30}[match.group(2)]
        start = int(now - days * 86400)
        return start, f"last {days} days (since {dt.datetime.fromtimestamp(start):%Y-%m-%d})"
    try:
        start = int(dt.datetime.strptime(text, "%Y-%m-%d").timestamp())
    except ValueError:
        raise excel.ExcelError("--since takes a span such as 30d, 12w, or 6m, a date (YYYY-MM-DD), or all.")
    return start, f"since {text}"


# --------------------------------------------------------------------------
# Collecting


def _inventory(app: Any) -> List[Dict[str, Any]]:
    from netops_platform.core import GlobalOptions
    return app.inventory_rows("", GlobalOptions(limit=1_000_000))


def _history(app: Any, since: int) -> List[Dict[str, Any]]:
    return app.state.rows(
        "SELECT * FROM inventory_history WHERE event_type='field_changed' AND created_at>=? "
        "AND field_name IN ('name','ip','mac','serial','model','platform') ORDER BY created_at, id", (since,))


def _first_seen(app: Any) -> Dict[Tuple[str, str, str], int]:
    seen: Dict[Tuple[str, str, str], int] = {}
    for row in app.state.rows("SELECT entity_type, identity_kind, identity_value, MIN(first_seen) AS first_seen "
                              "FROM inventory_identity_state WHERE first_seen>0 GROUP BY 1,2,3"):
        seen[(row["entity_type"], row["identity_kind"], str(row["identity_value"]).lower())] = int(row["first_seen"])
    return seen


def _first_seen_for(seen: Mapping[Tuple[str, str, str], int], entity: str, row: Mapping[str, Any]) -> Optional[dt.datetime]:
    keys = [("mac", real_mac(row.get("mac"))), ("ip", str(row.get("ip") or "")),
            ("name", str(row.get("name") or row.get("hostname") or "").lower())]
    found = [seen[(entity, kind, value.lower())] for kind, value in keys if value and (entity, kind, value.lower()) in seen]
    return _when(min(found)) if found else None


def transitions(history: Sequence[Mapping[str, Any]], moves: Sequence[Mapping[str, Any]],
                switch_ips: Iterable[str] = ()) -> List[Dict[str, Any]]:
    """The Transitions report's rows (see its rule)."""
    switch_ips = set(switch_ips)
    merged: Dict[Tuple[Any, ...], Dict[str, Any]] = {}
    for event in history:
        entity, field, kind = event["entity_type"], event["field_name"], event["identity_kind"]
        old, new = event.get("previous_value"), event.get("current_value")
        if entity == "ap" and is_switch_row(event):
            continue
        change, before, after = "", "", ""
        if field == "name" and kind in {"mac", "ip"}:
            before, after = _name(old), _name(new)
            if before and after and excel.name_key(before) != excel.name_key(after):
                change = "Renamed"
        elif field == "ip" and kind in {"mac", "name"}:
            before, after = excel.ip_key(old), excel.ip_key(new)
            if before and after and before != after:
                change = "IP changed"
        elif field == "mac" and kind in {"ip", "name"}:
            before, after = real_mac(old), real_mac(new)
            if before and after and before != after:
                change = "IP now on another device" if kind == "ip" else "MAC changed"
                before, after = excel.format_mac_as(before, "colon", False), excel.format_mac_as(after, "colon", False)
        elif field == "serial" and entity == "device" and kind in {"ip", "name", "mac"}:
            before, after = real_serial(old), real_serial(new)
            if before and after and before != after:
                change = "Serial changed"
        elif field == "model" and entity == "device":
            if switch_model_changed(old, new):
                change, before, after = "Model changed", aruba_model_name(old) or str(old).strip(), aruba_model_name(new) or str(new).strip()
        elif field == "platform" and entity == "ap":
            before, after = ap_model(old), ap_model(new)
            if before and after and before != after:
                change = "Model changed"
        if not change:
            continue
        key = (int(event["created_at"]) // 60, entity, change, before, after)
        row = merged.setdefault(key, {
            "When": _when(event["created_at"]), "Type": "Switch" if entity == "device" else "AP", "Change": change,
            "Device": _name(event.get("name")) or excel.ip_key(event.get("ip")) or "", "kept": [],
            "From": before, "To": after,
            "Switch / Port": " ".join(part for part in (str(event.get("switch_name") or ""), str(event.get("switch_port") or "")) if part),
        })
        label = {"mac": "MAC", "ip": "IP", "name": "name"}[kind]
        if label not in row["kept"]:
            row["kept"].append(label)
    for move in moves:
        if move.get("ap_ip") in switch_ips:
            continue
        before = f"{move.get('previous_switch_ip') or '?'} port {move.get('previous_switch_port') or '?'}"
        after = f"{move.get('current_switch_ip') or '?'} port {move.get('current_switch_port') or '?'}"
        key = (int(move["created_at"]) // 60, "ap", "Moved", before, after)
        merged.setdefault(key, {"When": _when(move["created_at"]), "Type": "AP", "Change": "Moved",
                                "Device": _name(move.get("ap_name")) or str(move.get("ap_ip") or ""), "kept": ["MAC"],
                                "From": before, "To": after, "Switch / Port": after})
    rows = []
    for row in merged.values():
        kept = row.pop("kept")
        row["Kept The Same"] = ", ".join(kept)
        rows.append(row)
    rows.sort(key=lambda row: (row["When"] or dt.datetime.min), reverse=True)
    return rows


def shared_serials(history: Sequence[Mapping[str, Any]]) -> set:
    """Serials recorded at more than one switch address: not any one box's own."""
    positions: Dict[str, set] = {}
    for event in history:
        if event["entity_type"] == "device" and event["field_name"] == "serial" and event["identity_kind"] == "ip":
            for value in (event.get("previous_value"), event.get("current_value")):
                serial = real_serial(value)
                if serial:
                    positions.setdefault(serial, set()).add(str(event["identity_value"]))
    return {serial for serial, seen in positions.items() if len(seen) > 1}


def replacements(history: Sequence[Mapping[str, Any]], entity: str, unreliable: Iterable[str] = ()) -> List[Dict[str, Any]]:
    """Hardware replacements for one entity type (see the history reports' rules).

    `unreliable` serials (see shared_serials) never count as a change.
    Returns events oldest first: position, ip, when, old/new model, serial, MAC.
    """
    unreliable = set(unreliable)
    positions = ("ip", "name") if entity == "device" else ("name",)
    scans: Dict[Tuple[str, str, int], Dict[str, Any]] = {}
    for event in history:
        if event["entity_type"] != entity or event["identity_kind"] not in positions:
            continue
        if entity == "ap" and is_switch_row(event):
            continue
        position = str(event["identity_value"] or "").lower()
        if event["identity_kind"] == "name" and placeholder_name(position):
            continue
        field, old, new = event["field_name"], event.get("previous_value"), event.get("current_value")
        bucket = int(event["created_at"]) // SAME_SCAN_SECONDS
        scan = scans.setdefault((event["identity_kind"], position, bucket), {
            "when": int(event["created_at"]), "name": _name(event.get("name")) or position, "ip": excel.ip_key(event.get("ip")),
            "new_model": str(event.get("model") or "") if entity == "device" else ap_model(event.get("platform")),
            "new_serial": real_serial(event.get("serial")), "new_mac": real_mac(event.get("mac")),
            "new_platform": str(event.get("platform") or ""), "changed": {},
        })
        if field == "serial" and real_serial(old) and real_serial(new) and real_serial(old) != real_serial(new) \
                and not {real_serial(old), real_serial(new)} & unreliable:
            scan["changed"]["serial"] = (real_serial(old), real_serial(new))
        elif field == "mac" and real_mac(old) and real_mac(new) and real_mac(old) != real_mac(new):
            scan["changed"]["mac"] = (real_mac(old), real_mac(new))
        elif field == "model" and entity == "device" and switch_model_changed(old, new):
            scan["changed"]["model"] = (aruba_model_name(old) or str(old).strip(), aruba_model_name(new) or str(new).strip())
        elif field == "platform" and entity == "ap":
            before, after = ap_model(old), ap_model(new)
            makers = (_maker(old), _maker(new))
            if before and after and before != after:
                scan["changed"]["model"] = (before, after)
            elif all(makers) and makers[0] != makers[1]:
                scan["changed"]["model"] = (before or makers[0], after or makers[1])
    events: Dict[Tuple[int, str], Dict[str, Any]] = {}
    for (_, position, bucket), scan in scans.items():
        changed = scan["changed"]
        if entity == "device":
            # The base MAC belongs to the chassis: with it unchanged, a new
            # serial is the same switch read another way.
            replaced = ("mac" in changed and ("serial" in changed or "model" in changed)) or \
                (not scan["new_mac"] and "serial" in changed and "model" in changed)
        else:
            replaced = "mac" in changed or "model" in changed
        if not replaced:
            continue
        old_serial, new_serial = changed.get("serial", ("", scan["new_serial"]))
        old_mac, new_mac = changed.get("mac", ("", scan["new_mac"]))
        old_model, new_model = changed.get("model", (scan["new_model"], scan["new_model"]))
        # The same swap seen from the IP and from the name is one event.
        identity = new_serial or new_mac or new_model
        events.setdefault((bucket, identity), {
            "when": scan["when"], "name": scan["name"], "ip": scan["ip"], "position": position,
            "old_model": old_model, "new_model": new_model, "old_serial": old_serial, "new_serial": new_serial,
            "old_mac": old_mac, "new_mac": new_mac,
        })
    ordered = sorted(events.values(), key=lambda event: event["when"])

    def hardware(event: Mapping[str, Any], side: str) -> str:
        return event[f"{side}_serial"] or event[f"{side}_mac"] or event[f"{side}_model"]
    reverted = set()
    for index, event in enumerate(ordered):
        for later in ordered[index + 1:]:
            if later["name"] != event["name"]:
                continue
            if hardware(later, "new") == hardware(event, "old"):
                reverted.update({id(event), id(later)})
            elif (hardware(later, "old"), hardware(later, "new")) == (hardware(event, "old"), hardware(event, "new")):
                # The same swap read again after the old box answered once more.
                reverted.add(id(later))
    return [event for event in ordered if id(event) not in reverted]


PORT_SWAP_HISTORY_CONFIDENCE = 75


def port_swap_events(swaps: Sequence[Mapping[str, Any]], current: Sequence[Mapping[str, Any]],
                     identities: Sequence[Mapping[str, Any]], entity: str,
                     switch_names: Mapping[str, str]) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """Neighbour-port swaps (identity.neighbor_port_transitions) as report rows.

    A swap where the new device is in the inventory and the old one is not
    is a replacement at that position, even when the name and the address
    both changed. The old device's hardware comes from Netrune's identity
    memory; the same MAC on both sides makes it a rename, not a replacement,
    and without a MAC on both sides it stays undecided (Transitions only).
    Returns (history events for this entity, Transitions rows).
    """
    present = {excel.name_key(row.get("name") or row.get("hostname")): row for row in current}
    latest: Dict[str, Mapping[str, Any]] = {}
    for row in identities:
        if row.get("entity_type") != entity:
            continue
        name = excel.name_key(row.get("name"))
        if name and int(row.get("last_seen") or 0) >= int(latest.get(name, {}).get("last_seen") or 0):
            latest[name] = row
    events: List[Dict[str, Any]] = []
    rows: List[Dict[str, Any]] = []
    for swap in swaps:
        old_name, new_name = excel.name_key(swap["old_name"]), excel.name_key(swap["new_name"])
        new = present.get(new_name)
        # The old device must be one of this kind that left: a port that saw
        # an AP and now sees a switch is a different job, not a replacement.
        if not new or old_name in present or old_name not in latest:
            continue
        old = latest[old_name]
        # An AP still on its factory name (APf814.ddc7.e19c) carries its MAC.
        old_mac = real_mac(old.get("mac")) or normalize_mac(re.sub(r"^AP", "", str(swap["old_name"]), flags=re.I))
        new_mac = real_mac(new.get("mac"))
        renamed = bool(old_mac and old_mac == new_mac)
        undecided = not (old_mac and new_mac)
        where = f"{switch_names.get(swap['local_ip'], swap['local_ip'])} {swap['local_port']}"
        evidence = f"neighbour port {where}; " + (
            "rest of that switch unchanged" if not swap["other_link_changes"] else f"{swap['other_link_changes']} other link change(s) there")
        rows.append({"When": _when(swap["when"]), "Type": "Switch" if entity == "device" else "AP",
                     "Change": "Renamed (same MAC)" if renamed else "Replaced or renamed (no MAC to compare)" if undecided
                     else "Replaced at neighbour port",
                     "Device": _host_label(swap["new_name"]), "Kept The Same": evidence,
                     "From": _host_label(swap["old_name"]), "To": _host_label(swap["new_name"]), "Switch / Port": where})
        if renamed or undecided or swap["confidence"] < PORT_SWAP_HISTORY_CONFIDENCE:
            continue
        model = str(old.get("model") or "") if entity == "device" else ap_model(old.get("platform"), old.get("model"))
        events.append({
            "when": swap["when"], "name": _host_label(swap["new_name"]), "ip": swap.get("new_ip") or str(new.get("ip") or ""),
            "position": where, "old_model": model if _meaningful_model(model) else _host_label(swap["old_name"]),
            "new_model": str(new.get("model") or "") if entity == "device" else ap_model(new.get("platform"), new.get("model")),
            "old_serial": real_serial(old.get("serial")), "new_serial": real_serial(new.get("serial")),
            "old_mac": old_mac, "new_mac": new_mac,
        })
    return events, rows


def _hardware_label(serial: str, mac: str, model: str) -> str:
    ident = serial or (excel.format_mac_as(mac, "colon", False) if mac else "")
    return f"{ident} ({model})" if ident and model else ident or model or "unknown hardware"


def history_rows(current: Sequence[Mapping[str, Any]], events: Sequence[Mapping[str, Any]], entity: str,
                 first_seen: Mapping[Tuple[str, str, str], int]) -> List[Dict[str, Any]]:
    """One row per device Netrune holds, then each piece of hardware it replaced, newest first."""
    by_key: Dict[str, List[Mapping[str, Any]]] = {}
    for event in events:
        by_key.setdefault(excel.name_key(event["name"]), []).append(event)
    rows: List[Dict[str, Any]] = []
    ordered = sorted(current, key=lambda row: (_host_label(row.get("name") or row.get("hostname") or "").lower(), str(row.get("ip") or "")))
    for device in ordered:
        name = _host_label(device.get("name") or device.get("hostname") or "") or str(device.get("ip") or "")
        swaps = sorted(by_key.get(excel.name_key(name), []), key=lambda event: event["when"], reverse=True)
        state = "Active" if str(device.get("status") or "").upper() == "ONLINE" else "Offline"
        latest = swaps[0] if swaps else None
        model = (device.get("model") if entity == "device" else ap_model(device.get("platform"), device.get("model"))) or ""
        base = {
            "name": name, "IP Address": device.get("ip") or "", "State": state,
            "Manufacturer": device.get("manufacturer") or _maker(device.get("platform"), device.get("model")),
            "Model": model if _meaningful_model(model) else "",
            "MAC": excel.format_mac_as(device.get("mac"), "colon", False),
            "First Seen": _first_seen_for(first_seen, entity, device), "Last Seen": _when(device.get("last_seen")),
            "Hardware Change": f"Replaced {_hardware_label(latest['old_serial'], latest['old_mac'], latest['old_model'])}" if latest else "",
            "Change Date": _when(latest["when"]) if latest else None,
        }
        if entity == "device":
            base["Serial Number"] = excel.device_value(device, "serial") or ""
        else:
            base.update({"Switch": device.get("switch_name") or "", "Switch Port": device.get("switch_port") or ""})
        rows.append(base)
        for event in swaps:
            old = {
                "name": name, "IP Address": event["ip"] or device.get("ip") or "", "State": "Replaced",
                "Manufacturer": _maker(event["old_model"], event["old_model"]) or base["Manufacturer"],
                "Model": event["old_model"], "MAC": excel.format_mac_as(event["old_mac"], "colon", False) if event["old_mac"] else "",
                "First Seen": None, "Last Seen": _when(event["when"]),
                "Hardware Change": f"Replaced by {_hardware_label(event['new_serial'], event['new_mac'], event['new_model'])}",
                "Change Date": _when(event["when"]),
            }
            if entity == "device":
                old["Serial Number"] = event["old_serial"]
            else:
                old.update({"Switch": base["Switch"], "Switch Port": base["Switch Port"]})
            rows.append(old)
    return rows


def collect(app: Any, since: int) -> Dict[str, Any]:
    inventory = _inventory(app)
    switches = [row for row in inventory if row.get("kind") == "device"]
    ap_rows = [row for row in inventory if row.get("kind") == "ap"]
    aps = [row for row in ap_rows if not is_switch_row(row)]
    # Switch history comes from the raw per-address observations, which no
    # reconciliation mistake can rewrite; AP history still from the merged
    # history (ADR 0008).
    device_history = app.state.device_observation_history(0)
    history = [row for row in _history(app, since) if row["entity_type"] != "device"] + \
        [row for row in device_history if int(row["created_at"]) >= since]
    # Which serials are shared across switches is a property of all history,
    # not of the reporting period.
    all_serial_history = [row for row in device_history if row["field_name"] == "serial"]
    moves = app.state.rows("SELECT * FROM ap_move_events WHERE created_at>=? ORDER BY created_at", (since,))
    first_seen = _first_seen(app)
    switch_ips = {row.get("ip") for row in switches} | {row.get("ip") for row in ap_rows if is_switch_row(row)}
    online = [row for row in aps if str(row.get("status") or "").upper() == "ONLINE"]
    # What the neighbour tables say: a port that saw one device now sees another.
    swaps = neighbor_port_transitions(app.state.rows("SELECT * FROM topology_events WHERE created_at>=?", (since,)))["swaps"]
    identities = app.state.rows("SELECT entity_type, name, mac, serial, model, platform, last_seen FROM inventory_identity_state WHERE identity_kind='name'")
    switch_names = {str(row.get("ip")): _host_label(row.get("hostname") or row.get("name") or "") for row in switches}
    switch_port_events, switch_port_rows = port_swap_events(swaps, switches, identities, "device", switch_names)
    ap_port_events, ap_port_rows = port_swap_events(swaps, aps, identities, "ap", switch_names)
    data = {
        "active-aps": [{
            "Name": _host_label(row.get("name") or "") or row.get("ip"), "IP Address": row.get("ip") or "",
            "MAC": excel.format_mac_as(row.get("mac"), "colon", False),
            "Manufacturer": row.get("manufacturer") or _maker(row.get("platform")),
            "Model": ap_model(row.get("platform"), row.get("model")),
            "Switch": row.get("switch_name") or "", "Switch Port": row.get("switch_port") or "",
            "Last Seen": _when(row.get("last_seen")),
        } for row in sorted(online, key=lambda row: str(row.get("name") or "").lower())],
        "active-switches": [{
            "Hostname": _host_label(row.get("hostname") or row.get("name") or "") or row.get("ip"),
            "IP Address": row.get("ip") or "", "Manufacturer": row.get("manufacturer") or "",
            "Model": excel.device_value(row, "model") or "", "Serial Number": excel.device_value(row, "serial") or "",
            "Base MAC": excel.format_mac_as(row.get("mac"), "colon", False), "Platform": row.get("platform") or "",
            "Last Seen": _when(row.get("last_seen")),
        } for row in sorted(switches, key=lambda row: _host_label(row.get("hostname") or row.get("name") or "").lower())
            if str(row.get("status") or "").upper() == "ONLINE"],
        "transitions": sorted(transitions(history, moves, switch_ips) + switch_port_rows + ap_port_rows,
                              key=lambda row: (row["When"] or dt.datetime.min), reverse=True),
        "switch-history": [{("Hostname" if key == "name" else "Base MAC" if key == "MAC" else key): value
                            for key, value in row.items()}
                           for row in history_rows(switches, replacements(history, "device", shared_serials(all_serial_history)) + switch_port_events,
                                                   "device", first_seen)],
        "ap-history": [{("Name" if key == "name" else key): value for key, value in row.items()}
                       for row in history_rows(aps, replacements(history, "ap") + ap_port_events, "ap", first_seen)],
    }
    data["_notes"] = {"switch_rows_in_ap_inventory": len(ap_rows) - len(aps)}
    return data


# --------------------------------------------------------------------------
# Writing


def build(data: Mapping[str, Any], path: Path, names: Sequence[str], *, period: str, generated: str = "") -> Dict[str, Any]:
    openpyxl = excel.require_openpyxl()
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    generated = generated or dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    wb = openpyxl.Workbook()
    info = wb.active
    info.title = "Report Info"
    info["A1"] = "Netrune reports"
    info["A1"].font = Font(bold=True, size=14, color="1F3864")
    info["A2"] = f"Generated {generated}; changes cover {period}."
    info.append([])
    info.append(["Sheet", "Rows", "What it lists"])
    for cell in info[4]:
        cell.font = Font(bold=True)
    counts = {}
    for name in names:
        report = REPORTS[name]
        rows = data.get(name, [])
        counts[name] = len(rows)
        info.append([report.sheet, len(rows), report.rule])
        ws = wb.create_sheet(report.sheet)
        ws.append(list(report.columns))
        for row in rows:
            ws.append([row.get(column) if row.get(column) not in ("", None) else None for column in report.columns])
        for col, column in enumerate(report.columns, start=1):
            values = [len(str(row.get(column) or "")) for row in rows[:500]]
            ws.column_dimensions[get_column_letter(col)].width = min(60, max(10, len(column) + 2, *(value + 2 for value in values)))
            if column in {"When", "Last Seen", "First Seen", "Change Date"}:
                for r in range(2, len(rows) + 2):
                    ws.cell(row=r, column=col).number_format = "yyyy-mm-dd hh:mm"
        last = max(2, len(rows) + 1)
        table = Table(displayName=re.sub(r"[^A-Za-z]", "", report.sheet), ref=f"A1:{get_column_letter(len(report.columns))}{last}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
        ws.freeze_panes = "B2"
    notes = data.get("_notes", {})
    if notes.get("switch_rows_in_ap_inventory") and any(name in names for name in ("active-aps", "ap-history")):
        info.append([])
        info.append([f"{notes['switch_rows_in_ap_inventory']} row(s) in Netrune's AP inventory are switches seen as "
                     "neighbours; the AP sheets leave them out."])
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 8
    info.column_dimensions["C"].width = 120
    for row in info.iter_rows(min_row=5, max_col=3):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    return {"sheets": [REPORTS[name].sheet for name in names], "rows": counts}
